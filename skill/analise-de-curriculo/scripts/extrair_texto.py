#!/usr/bin/env python3
"""Extrai o texto dos currículos (PDF/DOCX) e da vaga de uma pasta.

Uso:
    extrair_texto.py <pasta> --saida <dir> [--vaga "texto livre"]

Grava em <dir>:
    extracao.json   manifesto: vaga (fonte + texto) e currículos (id, arquivo, texto, legível, motivo)
    textos/<id>.txt texto de cada currículo legível (para leitura e para validar trechos citados)

Regras:
- nunca altera nem apaga nada na pasta de origem;
- ignora a subpasta _resultado (e qualquer pasta/arquivo que comece com "_" ou ".");
- vaga, por precedência: texto passado em --vaga; vaga.md na pasta; PDF/DOCX cujo nome contenha "vaga" ou "anuncio";
- arquivo sem texto aproveitável é marcado como ilegível, com o motivo, sem estimar o conteúdo.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

MIN_CHARS = 200  # menos que isso não dá para avaliar um currículo
EXT_OK = {".pdf", ".docx"}
EXT_IGNORADAS_AVISO = {".doc", ".odt", ".rtf", ".png", ".jpg", ".jpeg"}


def slug(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def texto_pdf(caminho: Path) -> str:
    from pypdf import PdfReader

    leitor = PdfReader(str(caminho))
    if leitor.is_encrypted:
        try:
            leitor.decrypt("")
        except Exception as e:  # noqa: BLE001
            raise ValueError("PDF protegido por senha") from e
    paginas = [(p.extract_text() or "") for p in leitor.pages]
    return "\n".join(paginas)


def texto_docx(caminho: Path) -> str:
    import docx

    doc = docx.Document(str(caminho))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            partes.append(" | ".join(c.text for c in linha.cells))
    return "\n".join(partes)


def extrair(caminho: Path) -> tuple[str, bool, str]:
    """Retorna (texto, legivel, motivo)."""
    try:
        if caminho.suffix.lower() == ".pdf":
            txt = texto_pdf(caminho)
        else:
            txt = texto_docx(caminho)
    except Exception as e:  # noqa: BLE001
        return "", False, f"não foi possível abrir o arquivo ({type(e).__name__}: {e})"
    txt = re.sub(r"&(\w{1,3});", r"&\1", txt)  # artefato "R&S;" de alguns geradores de PDF
    txt = re.sub(r"[ \t]+", " ", txt)
    txt = re.sub(r"\n\s*\n+", "\n", txt).strip()
    util = re.sub(r"\W", "", txt)
    if len(util) < MIN_CHARS:
        motivo = "sem texto extraível (provável PDF escaneado/imagem)" if len(util) == 0 else f"texto insuficiente ({len(util)} caracteres úteis)"
        return txt, False, motivo
    return txt, True, ""


def eh_vaga(nome: str) -> bool:
    n = slug(nome)
    return "vaga" in n or "anuncio" in n


def ignorado(p: Path, raiz: Path) -> bool:
    rel = p.relative_to(raiz)
    return any(parte.startswith(("_", ".")) for parte in rel.parts)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pasta")
    ap.add_argument("--saida", required=True, help="pasta onde gravar extracao.json e textos/")
    ap.add_argument("--vaga", default="", help="texto livre da vaga (tem precedência)")
    a = ap.parse_args(argv)

    raiz = Path(a.pasta).expanduser().resolve()
    if not raiz.is_dir():
        print(f"ERRO: pasta não encontrada: {raiz}", file=sys.stderr)
        return 2
    saida = Path(a.saida).expanduser().resolve()
    (saida / "textos").mkdir(parents=True, exist_ok=True)

    arquivos = sorted(p for p in raiz.rglob("*") if p.is_file() and not ignorado(p, raiz))
    candidatos_vaga = [p for p in arquivos if p.suffix.lower() in EXT_OK and eh_vaga(p.name)]

    vaga = {"fonte": None, "arquivo": None, "texto": ""}
    if a.vaga.strip():
        vaga = {"fonte": "texto do comando", "arquivo": None, "texto": a.vaga.strip()}
    elif (raiz / "vaga.md").is_file():
        vaga = {"fonte": "vaga.md", "arquivo": "vaga.md", "texto": (raiz / "vaga.md").read_text(encoding="utf-8").strip()}
    elif candidatos_vaga:
        p = candidatos_vaga[0]
        txt, ok, motivo = extrair(p)
        vaga = {"fonte": "arquivo da pasta", "arquivo": p.name, "texto": txt, **({} if ok or txt else {"erro": motivo})}

    curriculos, avisos, usados = [], [], set()
    for p in arquivos:
        ext = p.suffix.lower()
        if p.name == "vaga.md" or p in candidatos_vaga:
            continue
        if ext not in EXT_OK:
            if ext in EXT_IGNORADAS_AVISO:
                avisos.append(f"{p.name}: formato {ext} não suportado (converta para PDF ou DOCX)")
            continue
        base = slug(p.stem) or "curriculo"
        cid, n = base, 2
        while cid in usados:
            cid, n = f"{base}-{n}", n + 1
        usados.add(cid)
        txt, ok, motivo = extrair(p)
        item = {"id": cid, "arquivo": str(p.relative_to(raiz)), "legivel": ok, "motivo": motivo, "caracteres": len(txt)}
        if ok:
            (saida / "textos" / f"{cid}.txt").write_text(txt, encoding="utf-8")
            item["texto_arquivo"] = f"textos/{cid}.txt"
        curriculos.append(item)

    manifesto = {"pasta": str(raiz), "vaga": vaga, "curriculos": curriculos, "avisos": avisos}
    (saida / "extracao.json").write_text(json.dumps(manifesto, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Vaga: {vaga['fonte'] or 'NÃO ENCONTRADA — perguntar ao usuário'}" + (f" ({vaga['arquivo']})" if vaga.get("arquivo") else ""))
    print(f"Currículos: {len(curriculos)} | legíveis: {sum(c['legivel'] for c in curriculos)}")
    for c in curriculos:
        if not c["legivel"]:
            print(f"  NÃO AVALIADO: {c['arquivo']} — {c['motivo']}")
    for av in avisos:
        print(f"  aviso: {av}")
    print(f"Manifesto: {saida / 'extracao.json'}")
    return 0 if vaga["fonte"] else 3


if __name__ == "__main__":
    sys.exit(main())
