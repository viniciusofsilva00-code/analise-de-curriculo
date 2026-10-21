"""Extração de texto (PDF/DOCX), fonte da vaga e arquivos ilegíveis."""
import json
import shutil

from conftest import REFERENCIA_PASTA
from extrair_texto import main as extrair


def _rodar(pasta, saida, *extra):
    cod = extrair([str(pasta), "--saida", str(saida), *extra])
    return cod, json.loads((saida / "extracao.json").read_text(encoding="utf-8"))


def test_caso_de_referencia(tmp_path):
    antes = sorted(p.name for p in REFERENCIA_PASTA.iterdir())
    cod, m = _rodar(REFERENCIA_PASTA, tmp_path)
    assert cod == 0
    assert m["vaga"]["arquivo"] == "00_anuncio_vaga_analista_rh.pdf"
    assert "REQUISITOS OBRIGATÓRIOS" in m["vaga"]["texto"]
    assert len(m["curriculos"]) == 10 and all(c["legivel"] for c in m["curriculos"])
    texto = (tmp_path / m["curriculos"][0]["texto_arquivo"]).read_text(encoding="utf-8")
    assert "R&S;" not in texto  # artefato do PDF corrigido
    assert sorted(p.name for p in REFERENCIA_PASTA.iterdir()) == antes  # nada mudou na pasta


def test_ilegivel_docx_resultado_e_precedencia(tmp_path):
    import docx

    pasta = tmp_path / "cvs"
    (pasta / "_resultado" / "2026-01-01-x").mkdir(parents=True)
    shutil.copy(REFERENCIA_PASTA / "01_curriculo_juliana_ferreira_costa.pdf", pasta)
    shutil.copy(REFERENCIA_PASTA / "01_curriculo_juliana_ferreira_costa.pdf", pasta / "_resultado" / "2026-01-01-x" / "copia.pdf")
    shutil.copy(REFERENCIA_PASTA / "00_anuncio_vaga_analista_rh.pdf", pasta)
    (pasta / "quebrado.pdf").write_bytes(b"%PDF-1.4 isto nao e um pdf")
    d = docx.Document()
    for _ in range(20):
        d.add_paragraph("Analista de RH com experiência em recrutamento, seleção e treinamento.")
    d.save(pasta / "cv_word.docx")
    docx.Document().save(pasta / "cv_vazio.docx")
    (pasta / "foto.png").write_bytes(b"x")

    cod, m = _rodar(pasta, tmp_path / "s1")
    arquivos = {c["arquivo"]: c for c in m["curriculos"]}
    assert set(arquivos) == {"01_curriculo_juliana_ferreira_costa.pdf", "quebrado.pdf", "cv_word.docx", "cv_vazio.docx"}
    assert arquivos["cv_word.docx"]["legivel"]
    assert not arquivos["quebrado.pdf"]["legivel"] and arquivos["quebrado.pdf"]["motivo"]
    assert not arquivos["cv_vazio.docx"]["legivel"] and "sem texto" in arquivos["cv_vazio.docx"]["motivo"]
    assert m["vaga"]["arquivo"] == "00_anuncio_vaga_analista_rh.pdf"
    assert any("foto.png" in a for a in m["avisos"])

    (pasta / "vaga.md").write_text("Vaga: Analista de DP", encoding="utf-8")
    _, m = _rodar(pasta, tmp_path / "s2")
    assert m["vaga"]["fonte"] == "vaga.md"
    _, m = _rodar(pasta, tmp_path / "s3", "--vaga", "Analista de folha, 3 anos")
    assert m["vaga"]["fonte"] == "texto do comando" and m["vaga"]["texto"] == "Analista de folha, 3 anos"


def test_sem_vaga_retorna_codigo_3(tmp_path):
    pasta = tmp_path / "cvs"
    pasta.mkdir()
    shutil.copy(REFERENCIA_PASTA / "02_curriculo_marcos_vinicius_santos.pdf", pasta)
    cod, m = _rodar(pasta, tmp_path / "s")
    assert cod == 3 and m["vaga"]["fonte"] is None
